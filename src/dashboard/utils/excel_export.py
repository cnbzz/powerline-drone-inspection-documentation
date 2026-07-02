"""Excel export utility — formatlı, renkli .xlsx üretici.

:func:`export_faults_to_excel` filtrelenmiş arıza DataFrame'ini 3 sheet'li
(Arızalar / Risk Özeti / Uçuş Özeti) renkli bir Excel'e dönüştürür ve
``bytes`` döner — Streamlit ``st.download_button`` ile doğrudan kullanılır.

Beklenen df kolonları: fault_id, fault_id_short, timestamp, flight_label,
fault_type_label, equipment_label, risk_level, risk_level_tr, risk_score,
status, status_tr.
"""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RISK_FILL_COLORS = {
    'critical': 'CF222E',  # Kırmızı
    'high': 'BC4C00',      # Turuncu
    'medium': 'BF8700',    # Sarı
    'low': '1A7F37',       # Yeşil
}

STATUS_FILL_COLORS = {
    'new': '0969DA',             # Mavi
    'reviewed': 'BF8700',        # Sarı
    'resolved': '1A7F37',        # Yeşil
    'false_positive': '8C959F',  # Gri
}


def _apply_header_style(cell) -> None:
    """Excel başlık stili."""
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='1F2328', end_color='1F2328', fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = Border(
        left=Side(style='thin', color='888888'),
        right=Side(style='thin', color='888888'),
        top=Side(style='thin', color='888888'),
        bottom=Side(style='thin', color='888888'),
    )


def _apply_cell_style(cell, bg_color=None, font_color='1F2328', bold=False) -> None:
    """Veri hücresi stili."""
    cell.font = Font(name='Calibri', size=10, color=font_color, bold=bold)
    if bg_color:
        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type='solid')
    cell.alignment = Alignment(horizontal='left', vertical='center')
    cell.border = Border(
        left=Side(style='thin', color='D8DDE3'),
        right=Side(style='thin', color='D8DDE3'),
        top=Side(style='thin', color='D8DDE3'),
        bottom=Side(style='thin', color='D8DDE3'),
    )


def export_faults_to_excel(df: pd.DataFrame, filename: str | None = None) -> bytes:
    """Filtrelenmiş arıza listesini renkli, formatlı Excel'e dönüştür.

    Returns: Excel dosya içeriği (bytes) — st.download_button ile kullanılabilir.
    """
    if filename is None:
        filename = f"hhats_arizalar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    wb = Workbook()

    # === SHEET 1: Arıza Listesi ===
    ws = wb.active
    ws.title = "Arızalar"

    # Rapor başlığı
    ws['A1'] = "HHATS — Havai Hat Arıza Tespit Sistemi"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1F2328')
    ws.merge_cells('A1:H1')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws['A2'] = f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='656D76')
    ws.merge_cells('A2:H2')
    ws['A2'].alignment = Alignment(horizontal='center')

    ws['A3'] = f"Toplam Kayıt: {len(df)}"
    ws['A3'].font = Font(name='Calibri', size=10, color='656D76')
    ws.merge_cells('A3:H3')
    ws['A3'].alignment = Alignment(horizontal='center')

    ws.row_dimensions[4].height = 10

    # Tablo başlıkları (5. satır)
    headers = ['Fault ID', 'Tarih', 'Uçuş', 'Arıza Tipi', 'Ekipman', 'Risk', 'Skor', 'Durum']
    for col_idx, header in enumerate(headers, start=1):
        _apply_header_style(ws.cell(row=5, column=col_idx, value=header))
    ws.row_dimensions[5].height = 28

    # Veri satırları
    for row_idx, (_, row) in enumerate(df.iterrows(), start=6):
        ws.cell(row=row_idx, column=1, value=row['fault_id_short'])
        ws.cell(row=row_idx, column=2,
                value=pd.to_datetime(row['timestamp']).strftime('%d.%m.%Y %H:%M'))
        ws.cell(row=row_idx, column=3, value=row['flight_label'])
        ws.cell(row=row_idx, column=4, value=row['fault_type_label'])
        ws.cell(row=row_idx, column=5, value=row['equipment_label'])
        ws.cell(row=row_idx, column=6, value=row['risk_level_tr'])
        ws.cell(row=row_idx, column=7, value=round(float(row['risk_score']), 1))
        ws.cell(row=row_idx, column=8, value=row['status_tr'])

        # Risk hücresi renkli
        risk_color = RISK_FILL_COLORS.get(row['risk_level'], 'FFFFFF')
        _apply_cell_style(ws.cell(row=row_idx, column=6),
                          bg_color=risk_color, font_color='FFFFFF', bold=True)

        # Status hücresi renkli
        status_color = STATUS_FILL_COLORS.get(row['status'], 'FFFFFF')
        _apply_cell_style(ws.cell(row=row_idx, column=8),
                          bg_color=status_color, font_color='FFFFFF', bold=True)

        # Diğer hücreler
        for col_idx in [1, 2, 3, 4, 5, 7]:
            _apply_cell_style(ws.cell(row=row_idx, column=col_idx))

    # Kolon genişlikleri
    column_widths = {'A': 12, 'B': 18, 'C': 32, 'D': 18, 'E': 22, 'F': 12, 'G': 8, 'H': 14}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = 'A6'  # başlık satırını dondur

    # === SHEET 2: Risk Özeti ===
    ws2 = wb.create_sheet("Risk Özeti")
    ws2['A1'] = "Risk Seviyesi Dağılımı"
    ws2['A1'].font = Font(name='Calibri', size=14, bold=True)
    ws2.merge_cells('A1:C1')

    for col_idx, header in enumerate(['Risk Seviyesi', 'Adet', 'Yüzde'], start=1):
        _apply_header_style(ws2.cell(row=3, column=col_idx, value=header))

    risk_counts = df['risk_level_tr'].value_counts()
    total = len(df) or 1
    for row_idx, (risk, count) in enumerate(risk_counts.items(), start=4):
        ws2.cell(row=row_idx, column=1, value=risk)
        ws2.cell(row=row_idx, column=2, value=int(count))
        ws2.cell(row=row_idx, column=3, value=f"%{count / total * 100:.1f}")

        risk_key = {'Kritik': 'critical', 'Yüksek': 'high',
                    'Orta': 'medium', 'Düşük': 'low'}.get(risk, 'low')
        color = RISK_FILL_COLORS.get(risk_key, 'FFFFFF')
        _apply_cell_style(ws2.cell(row=row_idx, column=1),
                          bg_color=color, font_color='FFFFFF', bold=True)
        _apply_cell_style(ws2.cell(row=row_idx, column=2))
        _apply_cell_style(ws2.cell(row=row_idx, column=3))

    for col, width in [('A', 18), ('B', 12), ('C', 12)]:
        ws2.column_dimensions[col].width = width

    # === SHEET 3: Uçuş Özeti ===
    ws3 = wb.create_sheet("Uçuş Özeti")
    ws3['A1'] = "Uçuş Bazında Arıza Dağılımı"
    ws3['A1'].font = Font(name='Calibri', size=14, bold=True)
    ws3.merge_cells('A1:E1')

    for col_idx, header in enumerate(['Uçuş', 'Toplam', 'Kritik', 'Yüksek', 'Diğer'], start=1):
        _apply_header_style(ws3.cell(row=3, column=col_idx, value=header))

    flight_summary = df.groupby('flight_label').agg(
        total=('fault_id', 'count'),
        critical=('risk_level', lambda x: (x == 'critical').sum()),
        high=('risk_level', lambda x: (x == 'high').sum()),
    ).reset_index()
    flight_summary['other'] = (
        flight_summary['total'] - flight_summary['critical'] - flight_summary['high']
    )

    for row_idx, row in enumerate(flight_summary.itertuples(), start=4):
        ws3.cell(row=row_idx, column=1, value=row.flight_label)
        ws3.cell(row=row_idx, column=2, value=int(row.total))
        ws3.cell(row=row_idx, column=3, value=int(row.critical))
        ws3.cell(row=row_idx, column=4, value=int(row.high))
        ws3.cell(row=row_idx, column=5, value=int(row.other))
        for col_idx in range(1, 6):
            _apply_cell_style(ws3.cell(row=row_idx, column=col_idx))

    for col, width in [('A', 35), ('B', 12), ('C', 12), ('D', 12), ('E', 12)]:
        ws3.column_dimensions[col].width = width

    # Bytes olarak döndür (Streamlit download_button için)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
